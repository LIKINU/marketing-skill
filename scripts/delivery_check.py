#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""delivery_check.py — 交付形态核对（必交附件 ＋ 封面要素 ＋ 文件名 ＋ 占位符）

為什麼要有它（2026-09-17）：
    官方（校园赛／招标／政企申报都一个样）把这几条写成**硬性红线**，
    但在本仓库里**没有任何一处会去看封面或文件名** ——
    grep 全 scripts/ 目录，`队长姓名`／`封面`／文件名规范，一处命中都没有。
    后果：这类要求只能靠人记得；漏了不会有任何报错（「零件全绿、传动轴断了」的另一颗螺丝）。

它管四件事：
    ① **必交附件核对表（可勾选）** —— 官方要交的东西逐项核，缺哪件报哪件；
       `--init` 生成模板，人可以直接把 `"已交": true` 勾上（或让脚本按 glob 自动认）。
    ② **封面要素** —— 读 .docx 首页，核「封面只保留指定的那几项」（官方常写「仅保留项目名称＋团队名称」）。
    ③ **文件名规范** —— 如「项目名称-团队名称-队长姓名」；并**专门抓占位符没填**（`【队长姓名】`）。
    ④ **占位符残留扫描** —— `【填】`／`{FILL}`／TODO／待补……交付稿出现这些＝没写完。

用法：
    python scripts/delivery_check.py --dir 交付目录 --init                 # 生成可勾选的附件核对表
    python scripts/delivery_check.py --dir 交付目录                        # 核对
    python scripts/delivery_check.py --dir 交付目录 --manifest 核对表.json --cover-allow "项目名称,团队名称"
    python scripts/delivery_check.py --dir 交付目录 --name-pattern "^(?P<项目>.+)-(?P<团队>.+)-(?P<队长>.+)$"

退出码：0 全部通过；1 有硬项未通过；2 执行错误
"""
import argparse
import fnmatch
import io
import json
import os
import re
import sys

from _common import OK, NG, WARN, HINT, INFO   # noqa: E402  统一符号，不要在各自文件里重定义

# ── 内置的「通用必交附件」模板（没有 rules.json 时用它兜底）──
DEFAULT_MANIFEST = {
    "项目名": "",
    "团队名": "",
    "队长姓名": "",
    "封面要素": ["项目名称", "团队名称"],
    "文件名模式": r"^(?P<项目名称>.+?)-(?P<团队名称>.+?)-(?P<队长姓名>.+?)$",
    "必交附件": [
        {"名称": "项目报告 Word（.docx）", "glob": "*.docx", "必需": True, "已交": False},
        {"名称": "项目报告 PDF", "glob": "*.pdf", "必需": True, "已交": False},
        {"名称": "传播 Demo 视频（赛道二硬要求）", "glob": "*.mp4|*.mov|*.m4v",
         "必需": True, "已交": False},
        {"名称": "Demo 脚本文档／分镜", "glob": "*脚本*|*分镜*", "必需": True, "已交": False},
        {"名称": "原创承诺及版权授权说明（签字扫描件）",
         "glob": "*承诺*|*授权*", "必需": True, "已交": False},
        {"名称": "宣传合规性自查说明（单独提交）",
         "glob": "*自查*|*合规*", "必需": True, "已交": False},
        # ⚠️ 2026-09-17 补（R2 合规视角第 4 条）：官方红线要求单独提交《AI 使用说明》，
        #    而内置清单里**根本没有这一项** —— 缺件永远不会被发现。
        {"名称": "AI 使用说明（单独提交，须含工具名称与使用范围）",
         "glob": "*AI*|*人工智能*", "必需": True, "已交": False},
        {"名称": "设计作品效果图/主视觉", "glob": "*.png|*.jpg|*.jpeg|*.ai|*.psd",
         "必需": False, "已交": False},
    ],
}

# ── 占位符／未完成标记 ──
#   分兩級，理由：2026-09-17 實測，「占位」「待补」在營銷文稿裡是**正常詞**
#   （「被頭號對手長期占位的品類大詞」「心智占位」「一手結論留空待補，執行後回填」），
#   一律當硬錯誤會到處誤報 —— 誤報多了，人就開始忽略這張表。
#   硬錯誤只留「模板佔位符」這種一眼是機器的記號；散文詞降為警告。
PLACEHOLDER_HARD = ["【填】", "{FILL}", "（填）", "TODO", "TBD", "XXX", "???",
                    "【队长姓名】", "【團隊名稱】", "【项目名称】",
                    "【團隊名】", "（未给出处：", "（未給出處："]
PLACEHOLDER_WARN = ["待补", "待填", "待確認", "待确认", "待核实", "待核實"]

# ── 封面「不該出現」的額外字段（官方通常要求封面只保留極少數項）──
COVER_FORBIDDEN = ["学号", "学號", "指导教师", "指导老师", "指導教師", "姓名：",
                   "电话", "手機", "手机", "邮箱", "郵箱", "@", "赛道", "賽道",
                   "专业", "專業", "班级", "班級"]
# 單獨成段的日期／時間（官方若寫「僅保留項目名稱＋團隊名稱」，日期也算多餘）
COVER_DATE_RE = re.compile(r"^[0-9]{4}\s*[-年./]\s*[0-9]{1,2}\s*([-月./]\s*[0-9]{1,2}\s*日?)?$")


def _is_intermediate(fn):
    """中間產物／備份不參與「文件名規範」校驗 —— 它們本來就不該提交，
    但在工作目錄裡必然存在，混進清單只會把真問題淹掉。"""
    base = os.path.basename(fn)
    return bool(re.search(r"（备份）|\(备份\)|_v[1-9]\b|^plan_|plan_build|"
                          r"plan_fixed|_定稿|skeleton|备份|模板", base))


def load_manifest(path, d, rules_path=""):
    """manifest 优先；没有就用 rules.json 里的 delivery 段；再没有就用内置模板。"""
    if path and os.path.exists(path):
        m = json.loads(open(path, encoding="utf-8").read())
        base = json.loads(json.dumps(DEFAULT_MANIFEST))
        base.update(m)
        return base, path
    if rules_path and os.path.exists(rules_path):
        try:
            r = json.loads(open(rules_path, encoding="utf-8").read())
        except Exception:
            r = {}
        base = json.loads(json.dumps(DEFAULT_MANIFEST))
        # 只认「必交附件」这样的显式字段，不猜
        for k in ("必交附件", "必交附件清單", "附件"):
            if isinstance(r.get(k), list):
                base["必交附件"] = r[k]
                break
        dl = r.get("delivery") or {}
        for k in ("必交附件", "附件清单"):
            if isinstance(dl.get(k), list):
                base["必交附件"] = dl[k]
                break
        for k in ("项目名", "团队名", "队长姓名", "封面要素", "文件名模式"):
            if r.get(k):
                base[k] = r[k]
        return base, (rules_path + "（rules.json）")
    return json.loads(json.dumps(DEFAULT_MANIFEST)), "（内置模板）"


def collect_files(d, maxdepth=3):
    out = []
    root = os.path.abspath(d)
    for dirpath, dirnames, filenames in os.walk(root):
        rel = os.path.relpath(dirpath, root)
        if rel != "." and rel.count(os.sep) + 1 > maxdepth:
            dirnames[:] = []
            continue
        dirnames[:] = [x for x in dirnames if not x.startswith((".", "__"))]
        for f in filenames:
            if f.startswith("."):
                continue
            out.append(os.path.relpath(os.path.join(dirpath, f), root))
    return out


def magic_match(patterns, files):
    """glob 支持 `|` 分隔多模式；对**文件名**与**相对路径**都试一次。"""
    hits = []
    for pat in str(patterns).split("|"):
        pat = pat.strip()
        if not pat:
            continue
        for f in files:
            if fnmatch.fnmatch(os.path.basename(f), pat) or fnmatch.fnmatch(f, pat):
                hits.append(f)
    return sorted(set(hits))


def check_attachments(m, files):
    rows = []
    for it in m.get("必交附件", []):
        name = it.get("名称", "?")
        req = bool(it.get("必需", True))
        hit = magic_match(it.get("glob", ""), files) if it.get("glob") else []
        ticked = bool(it.get("已交"))
        rows.append({"名称": name, "必需": req, "命中": hit,
                     "已勾选": ticked, "通过": bool(hit) or ticked,
                     "备注": it.get("备注", "")})
    return rows


def docx_text(path):
    try:
        import docx
    except ImportError:
        return None, None
    d = docx.Document(path)
    paras = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    body = "\n".join(paras)
    for tb in d.tables:
        for r in tb.rows:
            body += "\n" + " | ".join(c.text for c in r.cells)
    return paras, body


def check_cover(paras, allow):
    """封面＝ .docx 開頭、目錄之前的少量段落。這裡取前 6 個非空段落做判斷。"""
    head = paras[:6]
    facts = []
    for p in head:
        for bad in COVER_FORBIDDEN:
            if bad in p:
                facts.append(f"封面出现不该有的字段「{bad}」：{p[:40]}")
        if COVER_DATE_RE.match(p.strip()):
            facts.append(f"封面出现独立的日期段落「{p.strip()}」"
                         f"（官方若要求封面仅保留项目名称+团队名称，日期属多余项）")
    if allow:
        for p in head[:len(allow) + 2]:
            if not any(a in p for a in allow) and len(p) > 30:
                facts.append(f"封面段落可能多余（既不属 {'/'.join(allow)}，又很长）：{p[:40]}")
    return head, facts


# ── 正文匿名扫描（R2 合规视角第 2 条）──
#   官方（校园赛／政企申报）：「正文不得出现团队成员姓名、学号、学校、指导教师，
#   **违反取消参赛资格**」。原实现只查封面 6 段（COVER_FORBIDDEN），正文零扫描。
#   难点：正文里出现「大学」多半是**别人家的大学**（产学研网络／对标案例），
#   所以必须带上下文豁免，否则一份完全合规的稿会被判违规。
ANON_WORDS = ["指导教师", "指导老师", "指導教師", "导师：", "導師：", "队长：", "隊長：",
              "组员：", "組員：", "组别：", "學號", "学号"]
ANON_SAFE = ["产学研", "產學研", "对标", "對標", "参考", "參考", "案例", "承办", "承辦",
             "主办", "主辦", "合作院校", "白皮书", "白皮書"]


def anon_scan(text):
    """回傳 [问题描述...]。"""
    facts = []
    for ln in text.split("\n"):
        safe = any(x in ln for x in ANON_SAFE) or "《" in ln
        for w in ANON_WORDS:
            if w in ln:
                facts.append(f"正文出现身份词「{w}」：{ln.strip()[:40]}")
        if not safe:
            for m in re.finditer(r"(?:大学|大學|学院|學院|高等专科)", ln):
                facts.append(f"正文出现院校名（非对标语境）：{ln.strip()[:40]}")
                break
        if "学号" in ln or "學號" in ln:
            if re.search(r"\d{6,14}", ln):
                facts.append(f"正文疑似出现学号数字：{ln.strip()[:40]}")
    return facts[:12]


def check_declaration(path):
    """声明件内容校验（R2 合规视角第 4 条）：
    原实现只做 glob 文件名匹配 —— **空文件也算「已交」**。AI 使用说明是官方红线的
    「单独提交件」，必须真的写清「工具名称 + 使用范围」。"""
    facts = []
    base = os.path.basename(path).lower()
    txt = ""
    if base.endswith(".docx"):
        try:
            import docx
            d = docx.Document(path)
            txt = "\n".join(p.text for p in d.paragraphs)
            for tb in d.tables:
                for r in tb.rows:
                    txt += "\n" + " | ".join(c.text for c in r.cells)
        except Exception as e:
            return [f"声明件读取失败（{type(e).__name__}）—— 无法校验内容，请人工看一遍"]
    elif base.endswith((".md", ".txt")):
        try:
            txt = io.open(path, encoding="utf-8").read()
        except Exception as e:
            return [f"声明件读取失败（{type(e).__name__}）"]
    else:
        return []
    if len(re.sub(r"\s", "", txt)) < 50:
        facts.append(f"声明件内容过短（{len(re.sub(chr(92)+'s','',txt))} 实字）—— 疑似空文件")
    if "AI" in os.path.basename(path) or "人工智能" in os.path.basename(path):
        has_tool = bool(re.search(r"ChatGPT|GPT|Claude|Kimi|豆包|文心|通义|Copilot|DeepSeek|"
                                  r"Midjourney|即梦|可灵|元宝|智谱|Gemini", txt))
        has_scope = bool(re.search(r"使用(的)?(范围|環節|环节|情况|情況)|用于|用於", txt))
        if not has_tool:
            facts.append("《AI 使用说明》没写**工具名称**（官方要求注明工具名称）")
        if not has_scope:
            facts.append("《AI 使用说明》没写**使用范围/环节**（官方要求注明使用范围）")
    return facts


def check_name(fname, pattern):
    stem = os.path.splitext(os.path.basename(fname))[0]
    facts = []
    for tok in ("【", "】", "{", "}", "<", ">", "TODO", "XX", "待填", "队长姓名"):
        if tok in os.path.basename(fname):
            facts.append(f"文件名里还有占位符「{tok}」没填")
            break
    if pattern:
        if not re.match(pattern, stem):
            facts.append(f"文件名不符合规范（期望三段式「项目名称-团队名称-队长姓名」，"
                         f"实际：「{stem}」）")
    return facts


def main():
    ap = argparse.ArgumentParser(description="交付形态核对（附件／封面／文件名／占位符）")
    ap.add_argument("--dir", required=True, help="交付目錄")
    ap.add_argument("--manifest", default="", help="附件核對表 JSON（--init 可生成）")
    ap.add_argument("--rules", default="", help="《任務規則表》JSON（其中若有必交附件／封面要求則採用）")
    ap.add_argument("--name-pattern", default="", help="文件名規範（正則，可含命名組）")
    ap.add_argument("--cover-allow", default="", help="封面允許的要素，逗號分隔（如 项目名称,团队名称）")
    ap.add_argument("--init", action="store_true", help="生成可勾選的附件核對表模板")
    ap.add_argument("--strict", action="store_true", help="有未通過項時以退出碼 1 結束（默認也返回 1）")
    a = ap.parse_args()

    if not os.path.isdir(a.dir):
        print(f"{NG} 找不到目錄：{a.dir}")
        sys.exit(2)

    man, src = load_manifest(a.manifest, a.dir, a.rules)
    if a.name_pattern:
        man["文件名模式"] = a.name_pattern
    if a.cover_allow:
        man["封面要素"] = [x.strip() for x in a.cover_allow.split(",") if x.strip()]

    mpath = a.manifest or os.path.join(a.dir, "附件核对表.json")
    if a.init:
        if os.path.exists(mpath):
            print(f"{WARN} 已存在，未覆蓋：{mpath}")
        else:
            open(mpath, "w", encoding="utf-8").write(
                json.dumps(man, ensure_ascii=False, indent=2))
            print(f"{OK} 已生成可勾選的附件核對表：{mpath}")
            print("   → 交了的把 \"已交\" 改成 true，或直接把檔案放進交付目錄（脚本按 glob 自動認）")
        man, src = load_manifest(mpath, a.dir, a.rules)

    files = collect_files(a.dir)
    fail = []

    print("=" * 68)
    print(f"交付形態核對 · delivery_check.py　目錄：{os.path.basename(os.path.abspath(a.dir))}")
    print(f"  核對表來源：{src}　｜　目錄內檔案 {len(files)} 個")
    print("=" * 68)

    # ① 必交附件
    print("\n【1】必交附件核對表")
    rows = check_attachments(man, files)
    if not rows:
        print(f"  {WARN} 核對表裡沒有附件項（可用 --init 生成模板）")
    for r in rows:
        mark = OK if r["通过"] else (NG if r["必需"] else WARN)
        how = ("已勾選" if r["已勾选"] else
               (f"找到：{r['命中'][0]}" if r["命中"] else "未找到"))
        print(f"  {mark} {'[必交]' if r['必需'] else '[选交]'} {r['名称']} —— {how}")
        if not r["通过"] and r["必需"]:
            fail.append(f"缺必交附件：{r['名称']}")

    # ② / ④ 讀 docx 做封面與佔位符
    docxes = [f for f in files if f.lower().endswith(".docx")
              and not f.startswith("~$") and "备份" not in f]
    print("\n【2】封面要素")
    cover_paras, cover_facts = [], []
    if not docxes:
        print(f"  {WARN} 目錄裡沒有 .docx，跳過封面與佔位符檢查")
    else:
        # ⚠️ 2026-09-17 改（R2 合規視角第 8 條）：原实现只查**最大的那一个 docx** ——
        #    多份交付件时（主报告 + 附件说明）其余文件完全不检，占位符与封面禁项可藏在第二份里。
        #    改成逐个检查。
        allow = man.get("封面要素") or []
        for target in docxes:
            paras, body = docx_text(os.path.join(a.dir, target))
            if paras is None:
                print(f"  {WARN} 缺 python-docx，跳過封面檢查（pip install python-docx）")
                break
            cover_paras, cover_facts = check_cover(paras, allow)
            print(f"  受檢檔案：{target}")
            print(f"    封面段落：" + " ｜ ".join(cover_paras[:4]))
            if cover_facts:
                for f in cover_facts:
                    print(f"    {NG} {f}")
                    fail.append(f)
            else:
                print(f"    {OK} 封面未發現多餘字段")
            # 正文匿名扫描（一票废标项）
            anon = anon_scan(body)
            if anon:
                for x in anon[:6]:
                    print(f"    {NG} 匿名红线：{x}")
                fail.append(f"{target} 正文匿名红线 {len(anon)} 处 —— 官方「违反取消参赛资格」")
            else:
                print(f"    {OK} 正文未見身份詞／院校名（非對標語境）")

    # ③ 文件名
    print("\n【3】文件名規範")
    pat = man.get("文件名模式") or DEFAULT_MANIFEST["文件名模式"]
    finals = [f for f in files
              if f.lower().endswith((".docx", ".pdf"))
              and not _is_intermediate(f)]
    skipped = [f for f in files if f.lower().endswith((".docx", ".pdf"))
               and _is_intermediate(f)]
    if skipped:
        print(f"  （跳過中間產物／備份 {len(skipped)} 個：{'、'.join(skipped[:3])}…）")
    if not finals:
        print(f"  {WARN} 沒有可提交的 .docx/.pdf 可查")
    for f in finals:
        facts = check_name(f, pat)
        if facts:
            for x in facts:
                print(f"  {NG} {f} —— {x}")
                fail.append(f"{f}：{x}")
        else:
            print(f"  {OK} {f}")

    # ④ 佔位符殘留
    print("\n【4】佔位符／未完成標記掃描")
    for target in docxes:
        _, body = docx_text(os.path.join(a.dir, target))
        if not body:
            continue
        hard = {t: body.count(t) for t in PLACEHOLDER_HARD if t in body}
        soft = {t: body.count(t) for t in PLACEHOLDER_WARN if t in body}
        if hard:
            print(f"  {NG} {target} 殘留模板佔位符 {len(hard)} 種："
                  + "、".join(f"{k}×{v}" for k, v in hard.items()))
            fail.append(f"{target} 殘留模板佔位符：{'、'.join(hard)}")
        if soft:
            print(f"  {WARN} {target} 出現 {len(soft)} 種「未完成」字樣（可能是主動聲明，"
                  f"也可能真沒寫完，需人工看一眼）："
                  + "、".join(f"{k}×{v}" for k, v in soft.items()))
        if not hard and not soft:
            print(f"  {OK} {target} 未見佔位符——稿子是真寫完的")

    # 声明件内容校验（R2-15b）
    print("\n【5】聲明件內容（不能只認檔名）")
    decl_files = [f for f in files
                  if re.search(r"AI|人工智能|承诺|承諾|授权|授權|自查|合规|合規", f)
                  and f.lower().endswith((".docx", ".md", ".txt"))]
    if not decl_files:
        print(f"  {WARN} 未見聲明件（AI 使用說明／原創承諾／合規自查）—— 官方要求單獨提交")
    for f in decl_files:
        df = check_declaration(os.path.join(a.dir, f))
        if df:
            for x in df:
                print(f"  {NG} {f} —— {x}")
                fail.append(f"{f}：{x}")
        else:
            print(f"  {OK} {f}　　内容非空且要素齐（工具名称/使用范围已核）")

    print("\n" + "=" * 68)
    if fail:
        print(f"{NG} 未通過 {len(fail)} 項：")
        for x in fail:
            print(f"   · {x}")
        print("\n→ 這些是**官方硬項**，不是建議；補齊後重跑本腳本。")
        sys.exit(1)
    print(f"{OK} 交付形態全部通過。")
    sys.exit(0)


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print(f"\n{WARN} 已中斷。")
        sys.exit(130)
    except Exception as e:
        print(f"\n{NG} delivery_check 執行出錯：{type(e).__name__}: {e}")
        print("→ 依協議 8：修正後重跑；環境問題就降級為人工逐項核對附件清單。")
        sys.exit(2)
